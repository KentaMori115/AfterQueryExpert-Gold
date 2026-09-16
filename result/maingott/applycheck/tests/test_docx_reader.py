"""DOCX extraction."""

from __future__ import annotations

from pathlib import Path

import pytest

from maingott_reel.errors import SourceError
from maingott_reel.models import BlockType
from maingott_reel.source.docx_reader import (
    heading_level,
    is_list_paragraph,
    normalize_text,
    read_docx,
)
from maingott_reel.utils.hashing import sha256_file


def test_normalize_text_collapses_whitespace_and_nbsp():
    assert normalize_text("  до 60   сек.  ") == "до 60 сек."


@pytest.mark.parametrize(
    ("style", "expected"),
    [
        ("Heading 1", 1),
        ("Heading 2", 2),
        ("Заголовок 3", 3),
        ("heading", 1),
        ("normal", None),
        ("List Bullet", None),
        (None, None),
    ],
)
def test_heading_level_detection(style: str | None, expected: int | None):
    assert heading_level(style) == expected


def test_missing_source_is_reported(tmp_path: Path):
    with pytest.raises(SourceError, match="not found"):
        read_docx(tmp_path / "absent.docx")


def test_non_docx_source_is_rejected(tmp_path: Path):
    path = tmp_path / "spec.txt"
    path.write_text("not a docx", encoding="utf-8")
    with pytest.raises(SourceError, match=r"must be a \.docx"):
        read_docx(path)


def test_corrupt_docx_is_reported(tmp_path: Path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"definitely not a zip archive")
    with pytest.raises(SourceError, match="Could not read"):
        read_docx(path)


def test_document_hash_matches_the_file(sample_docx: Path):
    assert read_docx(sample_docx).sha256 == sha256_file(sample_docx)


def test_blocks_are_numbered_in_document_order(sample_docx: Path):
    source = read_docx(sample_docx)
    assert [block.index for block in source.blocks] == list(range(source.block_count))
    assert [block.line for block in source.blocks] == list(range(1, source.block_count + 1))


def test_headings_lists_and_tables_are_typed(sample_docx: Path):
    source = read_docx(sample_docx)
    types = {block.block_type for block in source.blocks}
    assert BlockType.HEADING in types
    assert BlockType.PARAGRAPH in types
    assert BlockType.LIST_ITEM in types
    assert BlockType.TABLE_ROW in types


def test_blocks_carry_their_section_context(sample_docx: Path):
    source = read_docx(sample_docx)
    ai_blocks = [b for b in source.blocks if "RAG" in b.text]
    assert ai_blocks
    assert ai_blocks[0].section == "2. AI-контур и база знаний"

    target_rows = [b for b in source.blocks if b.block_type is BlockType.TABLE_ROW]
    assert target_rows[0].subsection == "1.1. Измеримый целевой эффект"


def test_table_rows_keep_their_coordinates(sample_docx: Path):
    source = read_docx(sample_docx)
    rows = [b for b in source.blocks if b.block_type is BlockType.TABLE_ROW]
    assert [row.row_index for row in rows] == [0, 1, 2]
    assert {row.table_index for row in rows} == {0}
    assert rows[1].text == "Время первого ответа в автоматических каналах — до 60 сек. для бота"


def test_counts_and_helpers(sample_docx: Path):
    source = read_docx(sample_docx)
    assert source.table_count == 1
    assert source.paragraph_count > 0
    assert source.block_count == len(source.blocks)
    assert "Содержание" in source.sections()
    assert source.text().splitlines()[0] == source.blocks[0].text


def test_pagination_is_absent_when_the_file_does_not_record_it(sample_docx: Path):
    source = read_docx(sample_docx)
    assert source.has_pagination is False
    assert all(block.page is None for block in source.blocks)


def test_extraction_is_deterministic(sample_docx: Path):
    first = read_docx(sample_docx)
    second = read_docx(sample_docx)
    assert first.model_dump(exclude={"extracted_at"}) == second.model_dump(exclude={"extracted_at"})


def test_real_specification_parses(real_spec: Path):
    source = read_docx(real_spec)
    assert source.block_count > 300
    assert source.table_count > 20
    assert any("MainGott" in block.text for block in source.blocks)


def test_numbered_paragraphs_are_list_items(tmp_path: Path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph("Пункт нумерованного списка спецификации.")
    paragraph._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>')
    )
    assert is_list_paragraph(paragraph) is True

    path = tmp_path / "numbered.docx"
    document.save(path)
    source = read_docx(path)
    assert source.blocks[0].block_type is BlockType.LIST_ITEM


def test_empty_table_rows_are_skipped(tmp_path: Path):
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Компонент"
    table.cell(0, 1).text = "Функции платформы MainGott и её контуров"
    # Second row deliberately left empty.
    path = tmp_path / "empty_row.docx"
    document.save(path)

    source = read_docx(path)
    rows = [b for b in source.blocks if b.block_type is BlockType.TABLE_ROW]
    assert [row.row_index for row in rows] == [0]
