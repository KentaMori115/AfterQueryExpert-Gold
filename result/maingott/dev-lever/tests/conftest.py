"""Shared test fixtures.

Tests never touch the real environment: settings are built explicitly and the
cached settings singleton is cleared around every test.
"""

from __future__ import annotations

from collections.abc import Callable, Iterator
from pathlib import Path

import pytest

from maingott_reel.config import Settings, reset_settings_cache
from maingott_reel.creative.draft import CreativePlanDraft, StoryboardDraft
from maingott_reel.models import CreativeBrief, FactRegistry, ScriptPlan, Storyboard
from maingott_reel.utils.run_context import RunContext

_ENV_VARS = (
    "OPENAI_API_KEY",
    "MAX_GENERATION_COST",
    "PRICING_FILE",
    "BRAND_REGISTRY",
    "VOICE_PROFILE",
    "RELEASE_VERSION",
    "REQUIRE_MUSIC",
    "OPENAI_TEXT_MODEL",
    "OPENAI_IMAGE_MODEL",
    "OPENAI_VIDEO_MODEL",
    "OPENAI_TTS_MODEL",
    "OUTPUT_ROOT",
    "INPUT_ROOT",
    "FFMPEG_BIN",
    "FFPROBE_BIN",
    "DEFAULT_REEL_DURATION",
    "SOURCE_DOCUMENT",
    "OPENAI_REASONING_EFFORT",
    "OPENAI_TEMPERATURE",
    "REEL_LANGUAGE",
    "ALLOW_TARGET_FACTS",
    "PROMPTS_ROOT",
    "LOG_LEVEL",
    "LOG_FORMAT",
)


@pytest.fixture(autouse=True)
def clean_environment(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Iterator[None]:
    """Isolate every test from the developer's .env and environment."""
    for name in _ENV_VARS:
        monkeypatch.delenv(name, raising=False)
    # Keep Rich output width stable so CLI assertions do not depend on the
    # terminal the tests happen to run in.
    monkeypatch.setenv("COLUMNS", "200")
    monkeypatch.chdir(tmp_path)
    reset_settings_cache()
    yield
    reset_settings_cache()


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    """Settings pointing at a temporary output tree."""
    return Settings(output_root=tmp_path / "output", input_root=tmp_path / "input")


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """A small DOCX that mirrors the structure of the real specification."""
    from docx import Document

    document = Document()
    document.add_heading("Содержание", level=1)
    document.add_paragraph("1. Назначение и бизнес-цели")

    document.add_heading("1. Назначение и бизнес-цели", level=1)
    document.add_paragraph(
        "Система представляет собой единую омниканальную платформу MainGott "
        "для привлечения, квалификации и продажи."
    )
    document.add_paragraph(
        "исключить повторный ввод контактных данных при переходе между каналами;",
        style="List Bullet",
    )
    document.add_paragraph("Слишком коротко")

    document.add_heading("1.1. Измеримый целевой эффект", level=2)
    metrics = document.add_table(rows=3, cols=2)
    metrics.cell(0, 0).text = "Показатель"
    metrics.cell(0, 1).text = "Целевое значение"
    metrics.cell(1, 0).text = "Время первого ответа в автоматических каналах"
    metrics.cell(1, 1).text = "до 60 сек. для бота"
    metrics.cell(2, 0).text = "Доля заявок с определённым источником обращения"
    metrics.cell(2, 1).text = "не менее 98%"

    document.add_heading("2. AI-контур и база знаний", level=1)
    document.add_paragraph(
        "AI-контур включает базу знаний, RAG, память диалога и guardrails "
        "для проверки ответов ассистента."
    )
    document.add_paragraph(
        "Bitrix24 является единым источником правды о клиенте, сделке и оплатах."
    )
    document.add_paragraph(
        "Bitrix24 является единым источником правды о клиенте, сделке и оплатах."
    )

    path = tmp_path / "spec.docx"
    document.save(path)
    return path


@pytest.fixture
def real_spec() -> Path:
    """The real MainGott specification, if it is present in the repository."""
    path = (
        Path(__file__).resolve().parents[1]
        / "input/source/MainGott_Technical_Specification_ru.docx"
    )
    if not path.is_file():
        pytest.skip("MainGott specification is not available")
    return path


@pytest.fixture
def registry() -> FactRegistry:
    """A small fact registry covering every claim status."""
    from datetime import UTC, datetime

    from maingott_reel.models import ClaimStatus, Fact, FactCategory

    facts = [
        Fact(
            id="F-001",
            statement="Единая омниканальная платформа MainGott для продаж и управления.",
            category=FactCategory.POSITIONING,
            source_lines=[2],
        ),
        Fact(
            id="F-002",
            statement="Сайт, Telegram, MAX, VK, YouTube, Instagram и FarPost работают как каналы.",
            category=FactCategory.CHANNEL,
            source_lines=[3],
        ),
        Fact(
            id="F-003",
            statement="Bitrix24 — единый источник правды о клиенте, сделке и коммуникациях.",
            category=FactCategory.INTEGRATION,
            source_lines=[4],
        ),
        Fact(
            id="F-004",
            statement="AI-контур включает базу знаний, RAG, память диалога и guardrails.",
            category=FactCategory.AI,
            source_lines=[5],
        ),
        Fact(
            id="F-005",
            statement="Сквозной путь клиента ведёт от первого контакта до оплаты и follow-up.",
            category=FactCategory.CUSTOMER_JOURNEY,
            source_lines=[6],
        ),
        Fact(
            id="F-006",
            statement="Аналитика собирает воронку, загрузку команды и денежный поток.",
            category=FactCategory.ANALYTICS,
            source_lines=[7],
        ),
        Fact(
            id="F-900",
            statement="Целевое время первого ответа — до 60 секунд в автоматических каналах.",
            category=FactCategory.CUSTOMER_JOURNEY,
            claim_status=ClaimStatus.TARGET,
            source_lines=[8],
        ),
        Fact(
            id="F-901",
            statement="Платформой уже пользуются тысячи компаний по всей стране сегодня.",
            category=FactCategory.OTHER,
            claim_status=ClaimStatus.UNSUPPORTED,
            source_lines=[9],
        ),
    ]
    return FactRegistry(
        source_sha256="a" * 64, created_at=datetime(2026, 8, 19, tzinfo=UTC), facts=facts
    )


@pytest.fixture
def draft_beats() -> list[dict[str, object]]:
    """Beat definitions for a valid eight-beat plan."""
    return [
        {
            "order": 1,
            "kind": "framing",
            "purpose": "Показать разрозненность каналов",
            "narration": "Клиенты приходят отовсюду.",
            "on_screen_text": "Разные каналы",
            "visual_direction": "Тёмная сцена, расходящиеся потоки",
            "estimated_seconds": 4.0,
            "claims": [],
        },
        {
            "order": 2,
            "kind": "framing",
            "purpose": "Потеря контекста",
            "narration": "Контекст теряется между ними.",
            "on_screen_text": "Потерянный контекст",
            "visual_direction": "Разрывы в потоке данных",
            "estimated_seconds": 4.0,
            "claims": [],
        },
        {
            "order": 3,
            "kind": "factual",
            "purpose": "Представить платформу",
            "narration": "MainGott соединяет их в одну платформу.",
            "on_screen_text": "MAINGOTT",
            "visual_direction": "Сборка единой системы",
            "estimated_seconds": 5.0,
            "claims": [
                {
                    "claim": "MainGott — единая омниканальная платформа продаж и управления.",
                    "kind": "factual",
                    "source_fact_ids": ["F-001"],
                }
            ],
        },
        {
            "order": 4,
            "kind": "factual",
            "purpose": "Каналы",
            "narration": "Сайт, мессенджеры и площадки — один контур.",
            "on_screen_text": "Один контур",
            "visual_direction": "Каналы сходятся в ядро",
            "estimated_seconds": 5.0,
            "claims": [
                {
                    "claim": "Сайт, Telegram, VK, YouTube, Instagram и FarPost — каналы платформы.",
                    "kind": "factual",
                    "source_fact_ids": ["F-002"],
                }
            ],
        },
        {
            "order": 5,
            "kind": "factual",
            "purpose": "AI понимает задачу",
            "narration": "AI понимает задачу клиента.",
            "on_screen_text": "AI понимает задачу",
            "visual_direction": "Спокойный интерфейс диалога",
            "estimated_seconds": 5.0,
            "claims": [
                {
                    "claim": "AI-контур включает базу знаний, RAG и guardrails.",
                    "kind": "factual",
                    "source_fact_ids": ["F-004"],
                }
            ],
        },
        {
            "order": 6,
            "kind": "factual",
            "purpose": "Путь клиента",
            "narration": "Один клиент, один контекст, один путь.",
            "on_screen_text": "Один путь",
            "visual_direction": "Линия пути клиента",
            "estimated_seconds": 6.0,
            "claims": [
                {
                    "claim": "Путь клиента ведёт от первого контакта до оплаты.",
                    "kind": "factual",
                    "source_fact_ids": ["F-005", "F-003"],
                }
            ],
        },
        {
            "order": 7,
            "kind": "factual",
            "purpose": "Аналитика",
            "narration": "Каждое действие становится данными.",
            "on_screen_text": "Данные для решений",
            "visual_direction": "Панель аналитики",
            "estimated_seconds": 5.0,
            "claims": [
                {
                    "claim": "Аналитика собирает воронку и денежный поток.",
                    "kind": "factual",
                    "source_fact_ids": ["F-006"],
                }
            ],
        },
        {
            "order": 8,
            "kind": "brand",
            "purpose": "Бренд-закрытие",
            "narration": "MainGott. Sales and Operations OS.",
            "on_screen_text": "MAINGOTT",
            "visual_direction": "Логотип на тёмном фоне",
            "estimated_seconds": 6.0,
            "claims": [
                {
                    "claim": "MainGott — единая платформа продаж и управления.",
                    "kind": "factual",
                    "source_fact_ids": ["F-001"],
                }
            ],
        },
    ]


@pytest.fixture
def make_draft(draft_beats: list[dict[str, object]]) -> Callable[..., CreativePlanDraft]:
    """Build a CreativePlanDraft, optionally overriding the beats."""

    def _make(
        beats: list[dict[str, object]] | None = None, **brief_overrides: object
    ) -> CreativePlanDraft:
        brief = {
            "objective": "Представить MainGott как единую систему продаж и операций.",
            "audience": "Владельцы и руководители бизнеса",
            "tone": "Спокойный, уверенный, профессиональный",
            "visual_direction": "Тёмная премиальная среда, чистая анимация данных",
            "core_message": "MainGott — Sales & Operations OS.",
            "supporting_messages": ["Один связанный путь клиента."],
            "cta": "MAINGOTT",
            "restrictions": [],
            "source_fact_ids": ["F-001", "F-002"],
        }
        brief.update(brief_overrides)
        return CreativePlanDraft.model_validate(
            {"brief": brief, "beats": beats if beats is not None else draft_beats}
        )

    return _make


@pytest.fixture
def analysed_cli_run(tmp_path: Path, sample_docx: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Run the analyze stage so plan commands have Phase 1 output to consume."""
    from maingott_reel.config import Settings
    from maingott_reel.source.analyzer import analyze

    monkeypatch.setenv("SOURCE_DOCUMENT", str(sample_docx))
    settings = Settings(source_document=sample_docx)
    return analyze(settings).run.root


@pytest.fixture
def script_plan(make_draft: Callable[..., CreativePlanDraft]) -> ScriptPlan:
    """A validated eight-beat plan, as the plan stage would have written it."""
    from datetime import UTC, datetime

    from maingott_reel.creative.script_generator import build_beats, build_plan
    from maingott_reel.models import (
        GenerationProvenance,
        Language,
        ValidationCheck,
        ValidationReport,
    )

    moment = datetime(2026, 8, 19, tzinfo=UTC)
    return build_plan(
        beats=build_beats(make_draft()),
        target_duration_seconds=40,
        language=Language.RU,
        validation=ValidationReport(
            created_at=moment, checks=[ValidationCheck(name="facts_exist", passed=True)]
        ),
        provenance=GenerationProvenance(
            generator_version="1.0",
            prompt_version="1/1",
            system_prompt_sha256="0" * 64,
            user_prompt_sha256="0" * 64,
            provider="fake",
            model="scripted-planner",
            generated_at=moment,
            source_sha256="a" * 64,
            offered_fact_ids=["F-001", "F-002", "F-003", "F-004", "F-005", "F-006"],
        ),
        created_at=moment,
    )


@pytest.fixture
def make_shots(script_plan: ScriptPlan) -> Callable[..., StoryboardDraft]:
    """Build a StoryboardDraft covering the plan, with optional overrides."""
    from maingott_reel.creative.draft import DraftShot

    house_prompt = (
        "Slow orbital camera move around a dark premium environment where flowing lines of "
        "light converge into a single calm core, polished glass surfaces, cinematic key "
        "light, professional mood, clean negative space across the lower third."
    )

    def _make(overrides: dict[str, dict[str, object]] | None = None) -> StoryboardDraft:
        shots = []
        for beat in script_plan.beats:
            fields: dict[str, object] = {
                "beat_id": beat.id,
                "visual_description": f"Тёмная сцена для {beat.id}.",
                "video_prompt": house_prompt,
                "transition": "cut",
            }
            fields.update((overrides or {}).get(beat.id, {}))
            shots.append(DraftShot.model_validate(fields))
        return StoryboardDraft(shots=shots)

    return _make


@pytest.fixture
def storyboard_fixture(
    script_plan: ScriptPlan, make_shots: Callable[..., StoryboardDraft]
) -> Storyboard:
    """A validated storyboard for the plan fixture."""
    from datetime import UTC, datetime

    from maingott_reel.creative.claims import validate_storyboard
    from maingott_reel.creative.storyboard_generator import build_scenes, build_storyboard
    from maingott_reel.models import GenerationProvenance, Language

    moment = datetime(2026, 8, 19, tzinfo=UTC)
    scenes = build_scenes(script_plan, make_shots(), 40)
    return build_storyboard(
        scenes=scenes,
        plan=script_plan,
        target_seconds=40,
        language=Language.RU,
        validation=validate_storyboard(scenes, script_plan, 40, created_at=moment),
        provenance=GenerationProvenance(
            generator_version="1.0",
            prompt_version="1/3",
            system_prompt_sha256="0" * 64,
            user_prompt_sha256="0" * 64,
            provider="fake",
            model="offline-planner",
            generated_at=moment,
            source_sha256="a" * 64,
        ),
        created_at=moment,
    )


@pytest.fixture
def storyboarded_run(
    settings: Settings,
    registry: FactRegistry,
    script_plan: ScriptPlan,
    storyboard_fixture: Storyboard,
) -> RunContext:
    """A run carrying a complete, valid Phase 1-3 result."""
    from maingott_reel.models import StageName
    from maingott_reel.utils.jsonio import write_model
    from maingott_reel.utils.manifest_store import load_or_create_manifest, save_manifest
    from maingott_reel.utils.run_context import create_run

    run = create_run(settings, run_id="run-1")
    write_model(run.facts_json, registry)
    write_model(run.script_json, script_plan)
    write_model(run.storyboard_json, storyboard_fixture)
    manifest = load_or_create_manifest(run)
    manifest.source_sha256 = registry.source_sha256
    manifest.record_stage(StageName.ANALYZE, completed_at=registry.created_at)
    manifest.record_stage(StageName.PLAN, completed_at=script_plan.created_at)
    manifest.record_stage(StageName.STORYBOARD, completed_at=storyboard_fixture.created_at)
    save_manifest(run, manifest)
    return run


@pytest.fixture(scope="session")
def ffmpeg_paths() -> tuple[str, str]:
    """Locate FFmpeg and ffprobe, or skip the test.

    The system binaries are preferred; the static build installed with the dev
    dependencies is the fallback so composition can be exercised in CI.
    """
    import shutil

    system = shutil.which("ffmpeg"), shutil.which("ffprobe")
    if all(system):
        return str(system[0]), str(system[1])
    try:
        from static_ffmpeg import run as static_run

        ffmpeg, ffprobe = static_run.get_or_fetch_platform_executables_else_raise()
        return str(ffmpeg), str(ffprobe)
    except Exception:  # noqa: BLE001 - any failure means we cannot compose here
        pytest.skip("FFmpeg is not available")


@pytest.fixture
def compose_settings(tmp_path: Path, ffmpeg_paths: tuple[str, str]) -> Settings:
    """Settings for composition tests: real FFmpeg, a smaller 9:16 canvas."""
    ffmpeg, ffprobe = ffmpeg_paths
    return Settings(
        output_root=tmp_path / "output",
        input_root=tmp_path / "input",
        brand_dir=tmp_path / "input" / "brand",
        ffmpeg_bin=ffmpeg,
        ffprobe_bin=ffprobe,
        video_width=540,
        video_height=960,
        video_fps=24,
        video_preset="ultrafast",
        video_crf=32,
    )


@pytest.fixture
def creative_brief() -> CreativeBrief:
    """The approved brief that accompanies the plan fixture."""
    return CreativeBrief(
        objective="Представить MainGott как единую систему продаж и операций.",
        audience="Владельцы и руководители бизнеса",
        target_duration_seconds=40,
        tone="Спокойный, уверенный, профессиональный",
        visual_direction="Тёмная премиальная среда, чистая анимация данных",
        core_message="MainGott — Sales & Operations OS.",
        cta="MAINGOTT",
        source_fact_ids=["F-001", "F-002"],
    )


@pytest.fixture
def composed_run(
    compose_settings: Settings,
    registry: FactRegistry,
    creative_brief: CreativeBrief,
    script_plan: ScriptPlan,
    storyboard_fixture: Storyboard,
) -> RunContext:
    """A run with a storyboard and real, decodable generated clips."""
    from maingott_reel.assets.manager import generate_assets
    from maingott_reel.models import StageName
    from maingott_reel.utils.jsonio import write_model
    from maingott_reel.utils.manifest_store import load_or_create_manifest, save_manifest
    from maingott_reel.utils.run_context import create_run

    run = create_run(compose_settings, run_id="compose-run")
    write_model(run.facts_json, registry)
    write_model(run.creative_brief_json, creative_brief)
    write_model(run.script_json, script_plan)
    write_model(run.storyboard_json, storyboard_fixture)
    from maingott_reel.models import SourceDocument

    write_model(
        run.source_json,
        SourceDocument(
            path=Path("input/source/spec.docx"),
            sha256=registry.source_sha256,
            extracted_at=registry.created_at,
            block_count=0,
        ),
    )

    manifest = load_or_create_manifest(run)
    manifest.source_sha256 = registry.source_sha256
    # Record what each stage would have recorded, so the audit sees a run that
    # looks like the pipeline produced it.
    manifest.files.update(
        {
            "source": run.source_json,
            "facts": run.facts_json,
            "creative_brief": run.creative_brief_json,
            "script": run.script_json,
            "storyboard": run.storyboard_json,
        }
    )
    manifest.record_stage(StageName.ANALYZE, completed_at=registry.created_at)
    manifest.record_stage(StageName.PLAN, completed_at=script_plan.created_at)
    manifest.record_stage(StageName.STORYBOARD, completed_at=storyboard_fixture.created_at)
    save_manifest(run, manifest)

    result = generate_assets(compose_settings, run_id=run.run_id, offline=True)
    assert result.complete, "the offline provider should produce every clip"
    return run


@pytest.fixture
def silent_voice_track(tmp_path: Path, ffmpeg_paths: tuple[str, str]) -> Path:
    """A short silent narration file, standing in for approved voice."""
    from maingott_reel.utils.ffmpeg import FFmpeg

    destination = tmp_path / "voice.m4a"
    FFmpeg(ffmpeg_paths[0]).encode(
        inputs=[Path("anullsrc=r=48000:cl=stereo")],
        output=destination,
        extra_input_args=[["-f", "lavfi"]],
        codec_args=["-t", "41", "-c:a", "aac", "-b:a", "96k"],
        label="test voice",
    )
    return destination


@pytest.fixture
def logo_file(tmp_path: Path) -> Path:
    """A stand-in brand mark, as an approved logo file would be."""
    from PIL import Image, ImageDraw

    path = tmp_path / "input" / "brand" / "logo.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGBA", (512, 128), (0, 0, 0, 0))
    ImageDraw.Draw(image).rectangle((16, 40, 496, 88), fill=(255, 255, 255, 255))
    image.save(path, "PNG")
    return path


# --- Phase 8: approved inputs -------------------------------------------------


def write_brand_registry(
    settings: Settings,
    logo: Path | None = None,
    music: Path | None = None,
    approved: bool = True,
) -> Path:
    """Write a brand registry describing the given assets."""
    from datetime import UTC, datetime

    from maingott_reel.models import AssetType, BrandAsset, BrandRegistry
    from maingott_reel.utils.hashing import sha256_file
    from maingott_reel.utils.jsonio import write_model

    assets = []
    if logo is not None:
        from PIL import Image

        with Image.open(logo) as image:
            width, height = image.size
            alpha = image.mode in ("RGBA", "LA")
        assets.append(
            BrandAsset(
                asset_id="maingott-logo",
                asset_type=AssetType.IMAGE,
                path=logo,
                sha256=sha256_file(logo),
                version="1.0",
                approved=approved,
                approved_by="Brand Owner" if approved else None,
                approved_at=datetime(2026, 8, 20, tzinfo=UTC) if approved else None,
                width=width,
                height=height,
                transparency=alpha,
            )
        )
    if music is not None:
        assets.append(
            BrandAsset(
                asset_id="maingott-music",
                asset_type=AssetType.MUSIC,
                path=music,
                sha256=sha256_file(music),
                approved=approved,
                approved_by="Brand Owner" if approved else None,
            )
        )
    path = settings.brand_registry_path
    write_model(path, BrandRegistry(updated_at=datetime(2026, 8, 20, tzinfo=UTC), assets=assets))
    return path


def write_voice_profile(
    settings: Settings,
    provider: str = "openai",
    model: str = "gpt-4o-mini-tts",
    voice: str = "marin",
    approved: bool = True,
) -> Path:
    """Write an approved (or unapproved) production voice profile."""
    from datetime import UTC, datetime

    from maingott_reel.models import Language, VoiceProfile
    from maingott_reel.utils.jsonio import write_model

    path = settings.voice_profile_path
    write_model(
        path,
        VoiceProfile(
            provider=provider,
            model=model,
            voice=voice,
            language=Language.RU,
            approved=approved,
            approved_by="Brand Owner" if approved else None,
            approved_at=datetime(2026, 8, 20, tzinfo=UTC) if approved else None,
            approval_notes="Listened to in Russian." if approved else None,
        ),
    )
    return path


def make_production_voice(settings: Settings, run: RunContext) -> None:
    """Generate narration offline, then record it as a real approved voice.

    No production speech model can be called here, so the *metadata* is made to
    look like one: the release gates read provenance, and this is the only way
    to exercise the production path without spending money.
    """
    from maingott_reel.audio.voice import generate_voice
    from maingott_reel.models import VoiceAsset
    from maingott_reel.utils.jsonio import read_model, write_model

    result = generate_voice(settings, run_id=run.run_id, offline=True)
    assert result.complete
    stored = read_model(run.voice_json, VoiceAsset)
    write_model(
        run.voice_json,
        stored.model_copy(
            update={
                "provider": "openai",
                "model": "gpt-4o-mini-tts",
                "voice": "marin",
                "development": False,
            }
        ),
    )


def make_production_assets(run: RunContext) -> None:
    """Record the run's offline clips as if a real video model had made them."""
    from maingott_reel.models import AssetCollection
    from maingott_reel.utils.jsonio import read_model, write_model

    collection = read_model(run.assets_json, AssetCollection)
    write_model(
        run.assets_json,
        collection.model_copy(
            update={
                "provider": "openai",
                "model": "sora-2",
                "assets": [
                    asset.model_copy(update={"provider": "openai", "model": "sora-2"})
                    for asset in collection.assets
                ],
            }
        ),
    )


@pytest.fixture
def production_settings(compose_settings: Settings, logo_file: Path) -> Settings:
    """Settings with an approved logo and an approved production voice."""
    settings = compose_settings.model_copy(update={"brand_logo": logo_file})
    write_brand_registry(settings, logo=logo_file)
    write_voice_profile(settings)
    return settings


@pytest.fixture
def production_run(production_settings: Settings, composed_run: RunContext) -> RunContext:
    """A run that looks like a real, approvable production run.

    The footage and the narration were still made offline — nothing here
    spends money — but they are recorded as a real provider's work so the
    release path can be exercised end to end.
    """
    from maingott_reel.audit.auditor import audit
    from maingott_reel.video.composition import compose

    make_production_voice(production_settings, composed_run)
    make_production_assets(composed_run)
    result = compose(production_settings, run_id=composed_run.run_id)
    assert result.passed, "the production fixture must compose cleanly"
    assert not result.development
    audit(production_settings, run_id=composed_run.run_id)
    return composed_run


@pytest.fixture
def reviewed_run(production_settings: Settings, production_run: RunContext) -> RunContext:
    """A production run whose human review is complete."""
    from maingott_reel.release import review as review_store
    from maingott_reel.release.inputs import fingerprint, load_inputs

    identity = fingerprint(production_settings, load_inputs(production_run))
    checklist, _ = review_store.current_checklist(production_run, identity.final_sha256)
    checklist = review_store.confirm(checklist, review_store.item_ids(), reviewer="Reviewer")
    review_store.save_checklist(production_run, checklist)
    return production_run


@pytest.fixture
def approved_run(production_settings: Settings, reviewed_run: RunContext) -> RunContext:
    """A production run a human has approved."""
    from maingott_reel.release import approval as approval_store
    from maingott_reel.release import review as review_store
    from maingott_reel.release.checker import release_check

    result = release_check(production_settings, run_id=reviewed_run.run_id)
    assert result.is_release_candidate, [gate.name for gate in result.blocking]
    approval_store.create(
        run=reviewed_run,
        fingerprint=result.fingerprint,
        approved_by="Release Manager",
        release_version="v1.0.0",
        review=review_store.load_checklist(reviewed_run),
        notes="Looked at on a phone.",
    )
    return reviewed_run
